"""Byg del1_theory.pptx til del 1: database design og CRUD.

Bygget på EK-skabelonen via kit.py. Del 1 bruger sin egen farvefamilie (use_part(1)).

Kør fra mappen sundhedsapp_postgres (luk præsentationen i PowerPoint først):

    uv run --with python-pptx slides/build_del1.py

eller inde fra mappen slides:

    uv run --with python-pptx build_del1.py

Et valgfrit argument gemmer præsentationen et andet sted: build_del1.py other.pptx
"""

import sys
from pathlib import Path

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from kit import (
    BAD, CODE, CORAL, DARK, GOOD, HEAD, INK, MINT, MONO, MUTED, PALE_TEXT, TEAL, WHITE,
    badge, card, code_box, hint_cards, inch, line, new_slide, notes, pt, rgb, roadmap, save, section, shape, table, text, title, title_slide, use_part, warning,
)

OUT = Path(__file__).parent.parent / "del1_database_crud" / "del1_theory.pptx"


# --- Intro -----------------------------------------------------------------------


def slide_title():
    title_slide('Lektion 5 · Del 1', 'Database design og CRUD',
                'Fra patient.yml til PostgreSQL med psycopg 3 og t-strings',
                ['ER diagram', 'CREATE TABLE', 'CRUD + t-strings', 'executemany()'],
                "Opsummering af lektion 4: API'et læser og skriver patient.yml, og /health_data kræver et bearer token. I dag flytter patienterne over i PostgreSQL.")


def slide_plan():
    s = new_slide()
    title(s, "Planen for del 1", "Trin 1–3 hjemme, trin 4–7 i undervisningen")
    columns = [
        ("Hjemme", TEAL, [
            ("P", "Installér PostgreSQL + pgAdmin, opret sundhedsapp"),
            ("P", "Læs theory.md og er_design_guide.md"),
            ("1", "Design ER diagrammet"),
            ("2", "CREATE TABLE i pgAdmin"),
            ("3", "Forbind fra Python"),
        ]),
        ("I undervisningen", CORAL, [
            ("4", "Create + read med t-strings"),
            ("5", "Update + delete"),
            ("6", "Lad databasen sige nej (CHECK)"),
            ("7", "API → executemany()"),
        ]),
    ]
    for c, (heading, color, steps) in enumerate(columns):
        x = 0.5 + c * 4.6
        text(s, x, 1.5, 4.3, 0.4, heading, font=HEAD, size=20, bold=True, color=color)
        for i, (label, body) in enumerate(steps):
            y = 2.0 + i * 0.6
            badge(s, x, y, label, color, size=0.42)
            text(s, x + 0.6, y, 3.8, 0.42, body, size=14, anchor=MSO_ANCHOR.MIDDLE)


def slide_why():
    s = new_slide()
    title(s, "Hvorfor forlade patient.yml?",
          "Lektion 4: læs hele filen, ændr en dict, skriv hele filen tilbage")
    items = [
        ("1", "To skrivere samtidig", "To requests gemmer i samme øjeblik, så én ændring går tabt."),
        ("2", "Ingen regler for data",
         "allergies: false, 'No' og Peanuts accepteres alle, og det gør age: -5 også."),
        ("3", "Det skalerer ikke", "For at finde én patient skal du læse alle patienter."),
    ]
    for i, (number, heading, body) in enumerate(items):
        x = 0.5 + i * 3.05
        card(s, x, 1.65, 2.85, 2.35)
        badge(s, x + 0.25, 1.9, number, CORAL)
        text(s, x + 0.25, 2.5, 2.4, 0.45, heading, font=HEAD, size=18, bold=True, color=DARK)
        text(s, x + 0.25, 2.95, 2.4, 0.95, body, size=14)
    text(s, 0.5, 4.4, 9, 0.5, "PostgreSQL håndterer mange brugere, håndhæver regler og finder rækker hurtigt.",
         size=17, bold=True, color=TEAL)


def slide_table():
    s = new_slide()
    title(s, "Tabeller, rækker og kolonner", "Én tabel = én slags ting. Én række = én patient.")
    header = ["patient_id", "first_name", "last_name", "age", "blood_type", "allergies"]
    rows = [
        ["1", "Kevin", "Holm", "34", "O+", "NULL"],
        ["2", "Lone", "Kellerman", "66", "A+", "Peanuts"],
        ["3", "Lars", "Larsson", "27", "B-", "Strawberry"],
    ]
    widths = [1.35, 1.45, 1.55, 0.95, 1.4, 2.3]
    grid = s.shapes.add_table(4, 6, inch(0.5), inch(1.6), inch(9), inch(1.68)).table
    for col, width in enumerate(widths):
        grid.columns[col].width = inch(width)
    for row_index, values in enumerate([header, *rows]):
        grid.rows[row_index].height = inch(0.42)
        for col, value in enumerate(values):
            cell = grid.cell(row_index, col)
            cell.margin_left = cell.margin_right = inch(0.06)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if row_index == 0:
                cell.fill.fore_color.rgb = rgb(CORAL if col == 0 else DARK)
                color, bold, italic = WHITE, True, False
            else:
                cell.fill.fore_color.rgb = rgb(MINT if row_index == 2 else WHITE)
                color = MUTED if value == "NULL" else INK
                bold, italic = col == 0, value == "NULL"
            run = cell.text_frame.paragraphs[0].add_run()
            run.text = value
            run.font.name = MONO
            run.font.size = pt(12)
            run.font.bold = bold
            run.font.italic = italic
            run.font.color.rgb = rgb(color)
    labels = [
        (CORAL, "Primary key", "Unikt id for én række. Genereres af databasen."),
        (DARK, "Kolonne", "Har et navn og en data type: INTEGER eller TEXT."),
        (TEAL, "Række", "Én patient. Fremhævet: Lone."),
    ]
    for i, (color, heading, body) in enumerate(labels):
        x = 0.5 + i * 3.05
        shape(s, MSO_SHAPE.OVAL, x, 3.6, 0.28, 0.28, color)
        text(s, x + 0.4, 3.53, 2.5, 0.4, heading, font=HEAD, size=16, bold=True, color=DARK)
        text(s, x + 0.4, 3.95, 2.5, 0.8, body, size=13)
    notes(s, "NULL er ikke en værdi, det betyder 'ingen værdi'. Vi bruger det til 'ingen kendte "
             "allergier' i stedet for false eller 'No'.")


# --- Afsnit 1: design af databasen -------------------------------------------------


def slide_er_blocks():
    s = new_slide()
    title(s, "ER diagram: byggestenene", "Tegn data, før du bygger dem")
    blocks = [
        ("Entity", "en ting → en tabel", "patient"),
        ("Attribute", "en oplysning → en kolonne", "age"),
        ("Primary key", "identificerer én række", "patient_id"),
        ("Foreign key", "peger på en anden PK", "hospital_id"),
        ("Relationship", "hvordan entities hænger sammen", "arbejder på"),
        ("Cardinality", "hvor mange på hver side", "1 : mange"),
    ]
    for i, (word, meaning, example) in enumerate(blocks):
        y = 1.5 + i * 0.55
        text(s, 0.5, y, 1.6, 0.45, word, font=HEAD, size=15, bold=True, color=DARK,
             anchor=MSO_ANCHOR.MIDDLE)
        text(s, 2.1, y, 2.1, 0.45, meaning, size=13, anchor=MSO_ANCHOR.MIDDLE)
        text(s, 4.2, y, 1.4, 0.45, example, font=MONO, size=11, color=TEAL,
             anchor=MSO_ANCHOR.MIDDLE)

    # Crow's foot-eksempel: HOSPITAL ||----o< DOCTOR
    text(s, 6.0, 1.5, 3.5, 0.35, "Crow's foot notation", font=HEAD, size=15, bold=True,
         color=DARK)
    for label, x in (("hospital", 6.0), ("doctor", 8.35)):
        shape(s, MSO_SHAPE.RECTANGLE, x, 2.0, 1.15, 0.5, DARK)
        text(s, x, 2.0, 1.15, 0.5, label, font=HEAD, size=13, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    y = 2.25
    line(s, 7.15, y, 8.35, y)
    line(s, 7.28, y - 0.12, 7.28, y + 0.12)          # || præcis én
    line(s, 7.38, y - 0.12, 7.38, y + 0.12)
    line(s, 8.1, y, 8.35, y - 0.13)                  # crow's foot: mange
    line(s, 8.1, y, 8.35, y + 0.13)
    shape(s, MSO_SHAPE.OVAL, 7.9, y - 0.07, 0.14, 0.14, WHITE, line=DARK, line_width=1.5)
    text(s, 6.0, 2.62, 3.5, 0.55,
         "Ét hospital har nul eller flere læger.\nHver læge arbejder på præcis ét hospital.",
         size=11, italic=True, color=MUTED)
    legend = [("||", "præcis én"), ("o|", "nul eller én"), ("|<", "én eller flere"),
              ("o<", "nul eller flere")]
    for i, (mark, meaning) in enumerate(legend):
        y = 3.35 + i * 0.36
        text(s, 6.0, y, 0.6, 0.32, mark, font=MONO, size=14, bold=True, color=CORAL)
        text(s, 6.6, y, 2.9, 0.32, meaning, size=13)
    text(s, 0.5, 4.75, 9, 0.3, "Del 1 har én entity. Relationships kommer i del 2.",
         size=13, italic=True, color=MUTED)


def slide_method():
    s = new_slide()
    title(s, "Design: en metode i 7 trin", "Hele gennemgangen står i er_design_guide.md")
    steps = [
        ("Indsaml", "Hvilke felter skal vi gemme?"),
        ("Entities", "Hvilke ting beskriver de?"),
        ("Attributes", "Tal eller tekst? Sortere eller regne?"),
        ("Primary key", "Hvad identificerer én række?"),
        ("Påkrævet?", "NULL eller NOT NULL?"),
        ("Regler", "Hvilke værdier giver ingen mening?"),
        ("Tjek", "Én værdi per celle? Gemt to gange?"),
    ]
    for i, (heading, question) in enumerate(steps):
        x = 0.5 + (i % 4) * 2.3
        y = 1.5 + (i // 4) * 1.6
        card(s, x, y, 2.1, 1.4)
        badge(s, x + 0.15, y + 0.15, str(i + 1), CORAL if i == 6 else TEAL, size=0.42)
        text(s, x + 0.15, y + 0.65, 1.85, 0.3, heading, font=HEAD, size=15, bold=True,
             color=DARK)
        text(s, x + 0.15, y + 0.95, 1.85, 0.4, question, size=11)
    text(s, 7.4, 3.1, 2.1, 1.4, "Tænk først, tegn bagefter. At ændre en tegning er billigt. "
         "At ændre en tabel fuld af data er ikke.", size=13, italic=True, color=TEAL,
         anchor=MSO_ANCHOR.MIDDLE)


def slide_method_1_3():
    s = new_slide()
    title(s, "Trin 1–3: felter, entity, data types", "Anvendt på patient.yml fra lektion 4")
    code_box(s, 0.5, 1.5, 3.3, 2.4,
             "patients:\n"
             "  id:\n"
             "    1:\n"
             "      age: 34\n"
             "      allergies: false\n"
             "      blood_type: O+\n"
             "      first_name: Kevin\n"
             "      last_name: Holm", size=11)
    text(s, 0.5, 4.05, 3.3, 1.0,
         [[("Entity: ", {"bold": True, "color": DARK}),
           ("alle felter beskriver én patient → tabellen ", {}),
           ("patient", {"font": MONO, "bold": True, "color": TEAL})],
          [("Navngiv tabeller med små bogstaver i ental.", {"italic": True, "color": MUTED})]],
         size=13, space_after=4)
    table(s, 4.1, 1.5, [1.35, 1.1, 2.95], [
        ["Felt", "Type", "Hvorfor"],
        ["first_name", "text", "fritekst"],
        ["last_name", "text", "fritekst"],
        ["age", "integer", "vi sorterer og sammenligner"],
        ["blood_type", "text", "tekst fra en fast liste"],
        ["allergies", "text", "tekst, efter rensning"],
    ], size=12, row_h=0.4, mono_cols=(0, 1))
    text(s, 4.1, 4.15, 5.4, 0.8,
         "Diskutér: alder ændrer sig hvert år. Et rigtigt system gemmer birth_date DATE og "
         "beregner alderen.", size=12, italic=True, color=MUTED)


def slide_method_4_6():
    s = new_slide()
    title(s, "Trin 4–6: key, påkrævet, regler")
    cards = [
        ("4", "Primary key", [
            "Skal være unik, aldrig NULL, må aldrig ændre sig.",
            "Navn? Der findes to Kevin Holm.",
            "CPR? Følsomme data.",
            "→ en surrogate key: patient_id, genereret af databasen.",
        ]),
        ("5", "Påkrævet?", [
            "Kan en patient eksistere uden?",
            "navne, age, blood_type → NOT NULL",
            "allergies → må være NULL",
            "Én betydning: NULL = ingen kendte allergier (ikke false, ikke 'No').",
        ]),
        ("6", "Regler", [
            "Hvilke værdier giver ingen mening?",
            "age: CHECK (age BETWEEN 0 AND 150)",
            "blood_type: CHECK (… IN ('A+', …, 'O-'))",
            "Databasen afviser dårlige data fra alle programmer.",
        ]),
    ]
    for i, (number, heading, lines_) in enumerate(cards):
        x = 0.5 + i * 3.05
        card(s, x, 1.25, 2.85, 3.1)
        badge(s, x + 0.2, 1.45, number, TEAL)
        text(s, x + 0.8, 1.45, 1.9, 0.45, heading, font=HEAD, size=18, bold=True, color=DARK,
             anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.2, 2.1, 2.5, 2.9, "\n".join(lines_), size=13, space_after=8)


def slide_method_7():
    s = new_slide()
    title(s, "Trin 7: tjek designet, og tegn så", "Resultatet for del 1")
    shape(s, MSO_SHAPE.RECTANGLE, 0.6, 1.5, 3.6, 0.5, DARK)
    text(s, 0.6, 1.5, 3.6, 0.5, "patient", font=HEAD, size=18, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    shape(s, MSO_SHAPE.RECTANGLE, 0.6, 2.0, 3.6, 2.7, WHITE, line=DARK, line_width=1.5)
    attributes = [
        ("PK", "patient_id", "integer"),
        ("", "first_name", "text"),
        ("", "last_name", "text"),
        ("", "age", "integer"),
        ("", "blood_type", "text"),
        ("", "allergies", "text · null"),
    ]
    for i, (key, name, data_type) in enumerate(attributes):
        y = 2.1 + i * 0.42
        middle = {"anchor": MSO_ANCHOR.MIDDLE, "font": MONO}
        text(s, 0.75, y, 0.45, 0.38, key, size=12, bold=True, color=CORAL, **middle)
        text(s, 1.25, y, 1.6, 0.38, name, size=13, **middle)
        text(s, 2.85, y, 1.3, 0.38, data_type, size=11, color=MUTED, **middle)
    hint_cards(s, 4.6, 1.5, 4.9, [
        ("Én værdi per celle?", "'Peanuts, Strawberry' i én celle kan ikke søges i. "
                                "Løsning: en allergy-tabel (ekstraopgave)."),
        ("Er noget gemt to gange?", "age og birth_date side om side kan være uenige."),
        ("Beskriver hver kolonne denne entity?",
         "hospital_name hører ikke til i patient. Den får sin egen tabel i del 2."),
    ], card_h=0.95, gap=0.12)


def slide_pgadmin_erd():
    s = new_slide()
    title(s, "Trin 1.7: tegn det i pgAdmin 9.18", "ERD Tool: fra tegning til CREATE TABLE")
    steps = [
        "Højreklik på sundhedsapp → ERD For Database",
        "Klik på Add table (tabelikon med +)",
        "Fanen General: Name = patient",
        "Fanen Columns: + for hver kolonne, vælg type",
        "Slå Not NULL? og Primary key? til",
        "patient_id → Constraints: IDENTITY, ALWAYS",
        "Save, derefter Generate SQL (SQL-ikon)",
        "Gem diagrammet som patient.pgerd",
    ]
    for i, step in enumerate(steps):
        x = 0.5 + (i // 4) * 4.6
        y = 1.5 + (i % 4) * 0.72
        badge(s, x, y, str(i + 1), TEAL, size=0.45)
        text(s, x + 0.6, y, 3.8, 0.45, step, size=14, anchor=MSO_ANCHOR.MIDDLE)
    warning(s, 4.55, "Bemærk: ", "ERD Tool laver ikke CHECK-regler. Tilføj dem i schema.sql "
                                "(trin 2.3).")


def slide_constraints():
    s = new_slide()
    title(s, "Trin 2: CREATE TABLE med constraints",
          "Hvis en regel brydes, afviser databasen statementet")
    code_box(s, 0.5, 1.5, 4.8, 3.4,
             "CREATE TABLE patient (\n"
             "  patient_id INTEGER\n"
             "    GENERATED ALWAYS AS IDENTITY\n"
             "    PRIMARY KEY,\n"
             "  first_name TEXT    NOT NULL,\n"
             "  last_name  TEXT    NOT NULL,\n"
             "  age        INTEGER NOT NULL\n"
             "    CHECK (age BETWEEN 0 AND 150),\n"
             "  blood_type TEXT    NOT NULL\n"
             "    CHECK (blood_type IN ('A+','A-',...)),\n"
             "  allergies  TEXT\n"
             ");", size=12)
    rules = [
        ("IDENTITY", "Nummereres 1, 2, 3… for dig. Erstatter len(new) + 1."),
        ("NOT NULL", "En værdi er påkrævet."),
        ("CHECK", "Værdien skal bestå en test: interval eller liste."),
        ("Ingen regel", "allergies må være NULL, dvs. ingen kendte allergier."),
    ]
    for i, (heading, body) in enumerate(rules):
        y = 1.5 + i * 0.87
        card(s, 5.55, y, 3.95, 0.77)
        text(s, 5.72, y + 0.08, 3.65, 0.3, heading, font=MONO, size=13, bold=True, color=CORAL)
        text(s, 5.72, y + 0.4, 3.65, 0.3, body, size=12)


def slide_step2_pitfalls():
    s = new_slide()
    title(s, "Trin 2: når SQL'en ikke vil køre", "Læs fejlen: den peger på linjen")
    table(s, 0.5, 1.5, [3.9, 2.8, 2.3], [
        ["Fejlen siger", "Årsag", "Løsning"],
        ["syntax error at or near \")\"", "komma efter sidste kolonne", "fjern sidste komma"],
        ["syntax error at or near \"last_name\"", "manglende komma på linjen før",
         "tilføj kommaet"],
        ["relation \"patient\" already exists", "tabellen blev oprettet tidligere",
         "behold DROP TABLE IF EXISTS først"],
        ["forkert INSERT fejlede ikke", "CHECK mangler", "tilføj CHECK, kør schema.sql igen"],
    ], size=12, row_h=0.5, mono_cols=(0,))
    text(s, 0.5, 4.2, 9, 0.8,
         [[("✅ Tjek: ", {"bold": True, "color": GOOD}),
           ("SELECT * FROM patient; viser én række med patient_id = 1 og allergies = [null]. "
            "De to forkerte INSERTs fra trin 2.6 skal fejle.", {})]], size=14)


# --- Afsnit 2: Python og psycopg 3 ---------------------------------------------------


def slide_crud():
    s = new_slide()
    title(s, "CRUD: de fire grundlæggende operationer")
    operations = [
        ("C", "Create", "INSERT INTO patient (first_name, age)\nVALUES ('Anna', 45)\nRETURNING patient_id;"),
        ("R", "Read", "SELECT * FROM patient\nWHERE patient_id = 1;"),
        ("U", "Update", "UPDATE patient SET age = 46\nWHERE patient_id = 1;"),
        ("D", "Delete", "DELETE FROM patient\nWHERE patient_id = 1;"),
    ]
    for i, (letter, name, sql) in enumerate(operations):
        x = 0.5 + (i % 2) * 4.6
        y = 1.2 + (i // 2) * 1.6
        card(s, x, y, 4.4, 1.4, CODE)
        badge(s, x + 0.2, y + 0.2, letter, TEAL if i < 2 else CORAL, 0.5)
        text(s, x + 0.85, y + 0.2, 3.3, 0.5, name, font=HEAD, size=18, bold=True, color=DARK,
             anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.85, y + 0.7, 3.45, 0.65, sql, font=MONO, size=11)
    warning(s, 4.5, "Advarsel: ", "UPDATE eller DELETE uden WHERE ændrer alle rækker. "
                                 "Skriv WHERE først.", h=0.6)


def slide_psycopg():
    s = new_slide()
    title(s, "psycopg 3: Python ↔ PostgreSQL",
          "connection → execute → cursor → fetch · with-blokken committer")
    flow = ["connect()", "execute(query)", "cursor", "fetchone() / fetchall()"]
    for i, step in enumerate(flow):
        x = 0.5 + i * 2.3
        highlight = i == 2
        card(s, x, 1.55, 2.0, 0.7, DARK if highlight else MINT)
        text(s, x, 1.55, 2.0, 0.7, step, font=MONO, size=12, bold=True,
             color=WHITE if highlight else DARK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        if i < len(flow) - 1:
            shape(s, MSO_SHAPE.RIGHT_ARROW, x + 2.04, 1.78, 0.22, 0.24, CORAL)
    code_box(s, 0.5, 2.55, 5.4, 1.75,
             "def read_patients() -> list[tuple]:\n"
             "    with get_connection() as conn:\n"
             "        cursor = conn.execute(\n"
             "            \"SELECT * FROM patient\"\n"
             "            \" ORDER BY patient_id\")\n"
             "        return cursor.fetchall()", size=11)
    facts = [
        ("fetchone()", "én tuple, eller None"),
        ("fetchall()", "liste af tuples"),
        ("rowcount", "rækker ændret af UPDATE/DELETE"),
        ("with … as conn", "commit ved succes, rollback ved fejl"),
    ]
    for i, (heading, body) in enumerate(facts):
        y = 2.55 + i * 0.6
        text(s, 6.15, y, 3.4, 0.28, heading, font=MONO, size=13, bold=True, color=TEAL)
        text(s, 6.15, y + 0.27, 3.4, 0.28, body, size=13)


def slide_injection():
    s = new_slide()
    title(s, "Byg aldrig SQL med f-strings", "SQL injection: data, der bliver til kode")
    columns = [
        (BAD, "✕", "f-string: farlig",
         "name = \"x' OR '1'='1\"\n"
         "conn.execute(f\"SELECT * FROM patient\"\n"
         "             f\" WHERE last_name = '{name}'\")",
         "f-stringen bliver først til almindelig tekst:\n… WHERE last_name = 'x' OR '1'='1'\n"
         "og alle patienter kommer tilbage."),
        (GOOD, "✓", "t-string: sikker",
         "name = \"x' OR '1'='1\"\n"
         "conn.execute(t\"SELECT * FROM patient\"\n"
         "             t\" WHERE last_name = {name}\")",
         "psycopg sender {name} som en separat parameter.\nDen er altid kun data, aldrig SQL.\n"
         "Intet kommer tilbage."),
    ]
    for i, (color, mark, heading, code, explanation) in enumerate(columns):
        x = 0.5 + i * 4.6
        badge(s, x, 1.45, mark, color)
        text(s, x + 0.6, 1.45, 3.8, 0.45, heading, font=HEAD, size=19, bold=True, color=color,
             anchor=MSO_ANCHOR.MIDDLE)
        code_box(s, x, 2.05, 4.4, 1.1, code, size=10)
        text(s, x, 3.3, 4.4, 1.0, explanation, size=13)
    text(s, 0.5, 4.6, 9, 0.5, "SQL injection er et af de mest almindelige sikkerhedshuller på "
         "nettet. Se OWASP i læselisten.", size=13, italic=True, color=MUTED)


def slide_tstrings():
    s = new_slide()
    title(s, "t-strings: template strings i Python 3.14",
          "De ligner f-strings, men tekst og værdier forbliver adskilt (PEP 750)")
    code_box(s, 0.5, 1.5, 5.3, 2.35,
             ">>> name = \"Holm\"\n"
             ">>> f\"WHERE last_name = {name}\"\n"
             "'WHERE last_name = Holm'\n"
             "\n"
             ">>> template = t\"WHERE last_name = {name}\"\n"
             ">>> template.strings\n"
             "('WHERE last_name = ', '')\n"
             ">>> template.values\n"
             "('Holm',)", size=11)
    hint_cards(s, 6.05, 1.5, 3.45, [
        ("f\"…\" → str", "Værdien bliver bagt ind i teksten med det samme."),
        ("t\"…\" → Template", "psycopg 3.3+ sender hver {value} som en parameter."),
    ], card_h=1.05, gap=0.2, body_size=12)
    text(s, 0.5, 4.05, 9, 0.9,
         [[("Hvorfor vi bruger dem: ", {"bold": True, "color": DARK}),
           ("værdien står, hvor den bruges, uden at tælle %s, uden tuple-rækkefølge at matche "
            "og uden (x,)-fælden.", {})],
          [("Kræver: ", {"bold": True, "color": DARK}),
           ("Python 3.14 og psycopg 3.3 eller nyere.", {})]], size=14, space_after=6)
    notes(s, "Live demo: kør 'uv run python', og skriv linjerne fra kodeboksen. "
             "Vis sql.as_string(t'...') fra psycopg for at se, hvad psycopg ville sende.")


def slide_specifiers():
    s = new_slide()
    title(s, "t-string format specifiers og %s", "Det meste af tiden behøver du ingen specifier")
    table(s, 0.5, 1.5, [0.9, 2.3, 1.8], [
        ["Spec", "Betydning", "Eksempel"],
        ["ingen", "værdi → parameter", "{age}"],
        [":i", "identifier: tabel/kolonne", "{column:i}"],
        [":l", "literal, flettet på klienten", "{payload:l}"],
        [":q", "et stykke SQL", "{where:q}"],
    ], size=12, row_h=0.42, mono_cols=(0, 2))
    text(s, 0.5, 3.75, 5.0, 1.3,
         "⚠️ En værdi kan aldrig være et kolonnenavn. ORDER BY {column} fejler. Brug {column:i}, "
         "og kun med navne fra en allow-list.", size=13, color=BAD)
    text(s, 5.8, 1.5, 3.7, 0.4, "Den klassiske måde: %s", font=HEAD, size=16, bold=True, color=DARK)
    code_box(s, 5.8, 1.95, 3.7, 1.0,
             "conn.execute(\n"
             "  \"… WHERE last_name = %s\",\n"
             "  (name,))", size=11)
    text(s, 5.8, 3.1, 3.7, 1.9,
         [[("Værdier skrives i en tuple, i samme rækkefølge som %s.", {})],
          [("Én værdi kræver et komma: ", {}), ("(name,)", {"font": MONO, "bold": True,
                                                            "color": GOOD})],
          [("Du skal bruge %s til executemany().", {"bold": True, "color": DARK})]],
         size=13, space_after=6)


# --- Afsnit 3: sådan løser du trin 4-7 ------------------------------------------------


def slide_step3():
    s = new_slide()
    title(s, "Trin 3: forbind fra Python", "uv run db.py → Forbundet! Server: PostgreSQL 17…")
    table(s, 0.5, 1.5, [4.4, 4.6], [
        ["Fejlen indeholder", "Løsning"],
        ["password authentication failed", "forkert password i DATABASE_URL i db.py"],
        ["database \"sundhedsapp\" does not exist", "opret databasen i pgAdmin (P2)"],
        ["Connection refused", "start PostgreSQL-servicen i Windows Services"],
        ["No module named 'psycopg'", "kør uv sync, og kør med uv run"],
        ["SyntaxError ved t\"…\"", "Python er ældre end 3.14: tjek python --version"],
    ], size=12, row_h=0.48, mono_cols=(0,))


def slide_step4():
    s = new_slide()
    title(s, "Trin 4: create_patient", "4.2 gør VALUES færdig · 4.3 kør og returnér id'et")
    code_box(s, 0.5, 1.5, 5.4, 2.75,
             "query = t\"\"\"\n"
             "    INSERT INTO patient (first_name, last_name,\n"
             "                         age, blood_type, allergies)\n"
             "    VALUES ({first_name}, {last_name},\n"
             "            {age}, {blood_type}, {allergies})\n"
             "    RETURNING patient_id\n"
             "\"\"\"\n"
             "with get_connection() as conn:\n"
             "    row = conn.execute(query).fetchone()\n"
             "return row[0]", size=11)
    hint_cards(s, 6.15, 1.5, 3.35, [
        ("Samme rækkefølge", "{values} følger rækkefølgen i kolonnelisten."),
        ("fetchone() → (5,)", "RETURNING giver én række med én værdi. row[0] er id'et."),
        ("4.4 test", "NotImplementedError: trin 4.5 betyder, at 4.2–4.3 virker."),
    ], card_h=1.0, gap=0.25)


def slide_step4_5():
    s = new_slide()
    title(s, "Trin 4.5 og 5: read, update, delete", "Ét mønster, tre funktioner")
    code_box(s, 0.5, 1.5, 5.6, 1.15,
             "with get_connection() as conn:\n"
             "    return conn.execute(\n"
             "        t\"SELECT * FROM patient\"\n"
             "        t\" WHERE patient_id = {patient_id}\").fetchone()",
             size=10)
    code_box(s, 0.5, 2.8, 5.6, 1.35,
             "with get_connection() as conn:\n"
             "    cursor = conn.execute(\n"
             "        t\"UPDATE patient SET age = {age}\"\n"
             "        t\" WHERE patient_id = {patient_id}\")\n"
             "return cursor.rowcount == 1", size=10)
    text(s, 0.5, 4.3, 5.6, 0.6,
         "5.2 delete_patient: det samme som update, med DELETE.",
         size=13, italic=True, color=TEAL)
    hint_cards(s, 6.35, 1.5, 3.15, [
        ("fetchone()", "en tuple, eller None når id'et ikke findes."),
        ("rowcount", "1 hvis en række blev ændret, 0 hvis id'et ikke fandtes."),
        ("Efter with", "Commit sker, når with-blokken slutter."),
    ], card_h=0.95, gap=0.15)


def slide_step6():
    s = new_slide()
    title(s, "Trin 6: lad databasen sige nej", "Fang fejlen, log den, og fortsæt")
    code_box(s, 0.5, 1.5, 5.6, 1.35,
             "try:\n"
             "    create_patient(\"Test\", \"Person\", 30, \"X+\")\n"
             "except CheckViolation as error:\n"
             "    logger.warning(\"Afvist af databasen: %s\",\n"
             "                   error.diag.message_primary)", size=11)
    text(s, 0.5, 3.3, 5.6, 0.35, "Output:", font=HEAD, size=14, bold=True, color=DARK)
    code_box(s, 0.5, 3.7, 9.0, 0.75,
             "Afvist af databasen: new row for relation \"patient\" violates check "
             "constraint \"patient_blood_type_check\"", size=10)
    hint_cards(s, 6.35, 1.5, 3.15, [
        ("Specifik exception", "Fang CheckViolation, ikke alle Exceptions."),
        ("Diskutér", "Skal reglen ligge i API'et, i databasen eller begge steder?"),
    ], card_h=0.95, gap=0.15)


def slide_executemany():
    s = new_slide()
    title(s, "executemany(): én SQL, mange rækker",
          "Samme statement, kørt én gang per tuple i en liste, i én transaction")
    code_box(s, 0.5, 1.55, 5.6, 1.75,
             "rows = [\n"
             "  (\"Kevin\", \"Holm\", 34, \"O+\", None),\n"
             "  (\"Lone\", \"Kellerman\", 66, \"A+\", \"Peanuts\"),\n"
             "]\n"
             "with get_connection() as conn, conn.cursor() as cur:\n"
             "    cur.executemany(INSERT_SQL, rows)", size=11)
    code_box(s, 0.5, 3.5, 5.6, 0.9,
             "cur.executemany(\"UPDATE … SET age = %s WHERE patient_id = %s\",\n"
             "                [(35, 1), (67, 2)])\n"
             "cur.executemany(\"DELETE … WHERE patient_id = %s\", [(3,), (4,)])", size=10)
    hint_cards(s, 6.35, 1.55, 3.15, [
        ("Hvorfor %s, ikke t\"…\"?", "En t-string holder værdierne for én række. executemany skal bruge tomme pladser."),
        ("På cursoren", "conn.cursor(), ikke conn."),
        ("Alt eller intet", "række 500 fejler → række 1–499 rulles tilbage."),
    ], card_h=0.95, gap=0.1, body_size=11)


def slide_api():
    s = new_slide()
    title(s, "Trin 7: fra API'et i lektion 4 til PostgreSQL",
          "Hent patienterne med dit bearer token, rens dem, og indsæt med executemany()")
    steps = [
        ("POST /token/1", "få \"Bearer eyJ…\""),
        ("GET /health_data", "header: Authorization"),
        ("to_row()", "false / 'No' → NULL"),
        ("executemany()", "INSERT alle rækker"),
    ]
    for i, (heading, body) in enumerate(steps):
        x = 0.5 + i * 2.3
        last = i == 3
        badge(s, x + 0.78, 1.6, str(i + 1), CORAL if i == 2 else TEAL)
        card(s, x, 2.2, 2.0, 1.2, DARK if last else MINT)
        text(s, x + 0.1, 2.3, 1.8, 0.45, heading, font=MONO, size=12, bold=True,
             color=WHITE if last else DARK, align=PP_ALIGN.CENTER)
        text(s, x + 0.1, 2.8, 1.8, 0.5, body, size=12,
             color=PALE_TEXT if last else INK, align=PP_ALIGN.CENTER)
    text(s, 0.5, 3.75, 9, 1.2,
         [[("7.1 ", {"bold": True, "color": CORAL}), ("Start appen fra lektion 4: flask run "
                                                        "(anden terminal)", {})],
          [("7.2 ", {"bold": True, "color": CORAL}), ("Prøv det i browseren: "
                                                        "http://127.0.0.1:5000/docs → Authorize", {})],
          [("7.3 ", {"bold": True, "color": CORAL}), ("Læs de færdige funktioner i bulk.py, "
                                                        "før du skriver dine egne", {})]],
         size=14, space_after=6)
    notes(s, "Dette er en migration: at flytte data fra ét lager til et andet og rense dem "
             "undervejs. Bemærk, at /add_patient fra lektion 4 gemmer 'bloodtype', mens YAML-filen "
             "bruger 'blood_type', og to_row() håndterer begge.")


def slide_step7():
    s = new_slide()
    title(s, "Trin 7: de tre funktioner, du skriver",
          "7.4–7.5 fetch_patients · 7.7 update_ages · 7.8 delete_patients")
    code_box(s, 0.5, 1.5, 9.0, 1.2,
             "headers = {\"Authorization\": token}        # token er allerede \"Bearer …\"\n"
             "response = requests.get(f\"{API_URL}/health_data\", headers=headers, timeout=5)\n"
             "response.raise_for_status()\n"
             "return response.json()[\"patients\"][\"id\"]", size=11)
    code_box(s, 0.5, 2.9, 5.6, 1.2,
             "sql = \"UPDATE patient SET age = %s WHERE patient_id = %s\"\n"
             "with get_connection() as conn, conn.cursor() as cur:\n"
             "    cur.executemany(sql, changes)\n"
             "    return cur.rowcount", size=10)
    code_box(s, 0.5, 4.25, 5.6, 0.6,
             "params = [(patient_id,) for patient_id in patient_ids]", size=10)
    hint_cards(s, 6.35, 2.9, 3.15, [
        ("(new_age, patient_id)", "Rækkefølgen i tuplen matcher rækkefølgen af %s."),
        ("[3, 4] → [(3,), (4,)]", "executemany skal bruge én tuple per række."),
    ], card_h=0.9, gap=0.15, body_size=11)


def slide_errors():
    s = new_slide()
    title(s, "Almindelige fejl og hvad de betyder")
    table(s, 0.5, 1.2, [4.6, 4.4], [
        ["Du ser", "Det betyder"],
        ["NotImplementedError: trin 4.5", "godt! Trinene før virker. Gå videre til 4.5"],
        ["TypeError: … not all arguments converted", "antallet af %s og værdier passer ikke"],
        ["UndefinedColumn: column \"bloodtype\"…", "stavefejl i et kolonnenavn: sammenlign med schema.sql"],
        ["CheckViolation", "data bryder en CHECK-regel, som tilsigtet"],
        ["NotNullViolation", "en påkrævet værdi mangler (None)"],
        ["requests ConnectionError", "appen fra lektion 4 kører ikke (7.1)"],
        ["401 UNAUTHORIZED", "token mangler i Authorization-headeren"],
    ], size=12, row_h=0.46, mono_cols=(0,))


# --- Afslutning ------------------------------------------------------------------------


def slide_links():
    s = new_slide()
    title(s, "Læs mere", "Alle links står også nederst i theory.md")
    groups = [
        ("PostgreSQL", [
            ("Officiel tutorial", "https://www.postgresql.org/docs/current/tutorial.html"),
            ("CREATE TABLE", "https://www.postgresql.org/docs/current/sql-createtable.html"),
            ("Constraints", "https://www.postgresql.org/docs/current/ddl-constraints.html"),
            ("Identity columns",
             "https://www.postgresql.org/docs/current/ddl-identity-columns.html"),
            ("INSERT, UPDATE, DELETE", "https://www.postgresql.org/docs/current/dml.html"),
            ("W3Schools SQL", "https://www.w3schools.com/sql/"),
        ]),
        ("psycopg 3 og Python", [
            ("t-string queries", "https://www.psycopg.org/psycopg3/docs/basic/tstrings.html"),
            ("Passing parameters (%s)",
             "https://www.psycopg.org/psycopg3/docs/basic/params.html"),
            ("Basic usage", "https://www.psycopg.org/psycopg3/docs/basic/usage.html"),
            ("Transactions", "https://www.psycopg.org/psycopg3/docs/basic/transactions.html"),
            ("PEP 750: Template strings", "https://peps.python.org/pep-0750/"),
            ("OWASP: SQL injection", "https://owasp.org/www-community/attacks/SQL_Injection"),
        ]),
        ("ER diagrammer", [
            ("er_design_guide.md", None),
            ("pgAdmin ERD Tool", "https://www.pgadmin.org/docs/pgadmin4/latest/erd_tool.html"),
            ("Lucidchart: ER diagrams", "https://www.lucidchart.com/pages/er-diagrams"),
            ("Mermaid ER syntax",
             "https://mermaid.js.org/syntax/entityRelationshipDiagram.html"),
        ]),
    ]
    for i, (heading, links) in enumerate(groups):
        x = 0.5 + i * 3.05
        card(s, x, 1.45, 2.85, 3.0)
        text(s, x + 0.2, 1.6, 2.5, 0.35, heading, font=HEAD, size=16, bold=True, color=DARK)
        lines_ = []
        for label, url in links:
            opts = {"color": TEAL, "link": url} if url else {"color": INK, "bold": True}
            lines_.append([(label, opts)])
        text(s, x + 0.2, 2.05, 2.5, 3.0, lines_, size=13, space_after=7)


def slide_roadmap():
    roadmap('Din tur: exercise.md, trin 4–7', [
        ('4', 'Create + read', 't-strings', 10),
        ('5', 'Update + delete', 'rowcount', 5),
        ('6', 'CHECK-fejl', 'try / except', 3),
        ('7', 'executemany()', 'API + token', 12),
    ], [
        [('Sidder du fast? ', {"bold": True, "color": DARK}),
         ('1) læs fejlen nedefra og op  2) læs TODO-hintet  3) find '
          'slidet for trinnet  4) kig i solution/', {})],
        [('Færdig før tid? ', {"bold": True, "color": DARK}),
         ('Ekstraopgaver: SQL injection-test, ORDER BY {column:i}, en '
          'allergy-tabel.', {})],
    ])


def main() -> None:
    use_part(1)
    slide_title()
    slide_plan()
    slide_why()
    slide_table()
    section(1, "Design af databasen", "ER diagram → CREATE TABLE (trin 1–2)")
    slide_er_blocks()
    slide_method()
    slide_method_1_3()
    slide_method_4_6()
    slide_method_7()
    slide_pgadmin_erd()
    slide_constraints()
    slide_step2_pitfalls()
    section(2, "Python og psycopg 3", "CRUD, t-strings og hvorfor f-strings er farlige")
    slide_crud()
    slide_psycopg()
    slide_injection()
    slide_tstrings()
    slide_specifiers()
    slide_executemany()
    section(3, "Sådan løser du trin 3–7", "Mønstre, hints og de tjek, du skal kigge efter")
    slide_step3()
    slide_step4()
    slide_step4_5()
    slide_step6()
    slide_api()
    slide_step7()
    slide_errors()
    slide_links()
    slide_roadmap()
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else OUT
    try:
        save(out, "Del 1 - Database design og CRUD")
    except PermissionError:
        sys.exit(f"Kan ikke skrive {out}. Luk den i PowerPoint, og kør igen.")
    print(f"skrev {out}")


if __name__ == "__main__":
    main()
