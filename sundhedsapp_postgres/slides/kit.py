"""Fælles byggesten til slides for del 1-4, bygget på EK's PowerPoint-skabelon.

Skabelonen (ek_powerpoint_skabelon.potx) har fem farvefamilier med hver sine layouts.
Hver del får sin egen familie via use_part(). Indholdet står på et hvidt panel oven på
familiens baggrund, så tabeller og kode er lette at læse.

Alle mål i builders er i "gamle" tommer på et 10 x 5,625 lærred. kit skalerer dem (og
skriftstørrelser) til skabelonens 13,33 x 7,5, så forholdene er de samme.
"""

import io
import zipfile
from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

TEMPLATE = Path(__file__).parent / "ek_powerpoint_skabelon.potx"
SCALE = 4 / 3   # 10 x 5,625 -> 13,333 x 7,5

# Farveroller. rgb() slår dem op i den aktive del's tema. Hex-strenge virker også direkte.
DARK = "dark"            # overskrifter, tabelhoveder, mørke kort
TEAL = "accent"          # badges og accenttekst
MINT = "light"           # lyse kort
CORAL = "highlight"      # fremhævning
INK = "ink"              # brødtekst
MUTED = "muted"          # noter
WHITE = "white"
CODE = "code"            # kodebokse
GOOD = "good"
BAD = "bad"
WARN_BG = "warn_bg"
PALE_TEXT = "pale_text"  # lys tekst på DARK-fyld
PALE_TEAL = "pale_text"
CARD_ON_DARK = "dark"

HEAD = "Aptos Serif"     # skabelonens titelfont
# Slides blev designet med Cambria. Aptos Serif er ca. 10 % bredere, så overskrifter i
# kort og bokse skaleres lidt ned for at passe i de samme bredder.
HEAD_SIZE_FACTOR = 0.9
BODY = "Calibri"         # skabelonens brødtekst
MONO = "Courier New"


@dataclass(frozen=True)
class Family:
    """Layouts i én af skabelonens farvefamilier."""

    title: str
    section: str
    content: str


FAMILIES = {
    "creme": Family("10_Titelslide", "15_Titelslide", "To indholdsobjekter"),
    "laks": Family("11_Titelslide", "16_Titelslide", "1_To indholdsobjekter"),
    "grøn": Family("12_Titelslide", "17_Titelslide", "2_To indholdsobjekter"),
    "rød": Family("13_Titelslide", "18_Titelslide", "3_To indholdsobjekter"),
    "blå": Family("14_Titelslide", "19_Titelslide", "4_To indholdsobjekter"),
}

COMMON = {
    "ink": "333333", "muted": "5F6B73", "white": "FFFFFF", "code": "F2F3F4",
    "good": "2E8B57", "bad": "C0392B", "warn_bg": "FDE8E2",
}

# Én farve per del. on_bg er tekst direkte på familiens baggrund.
PARTS = {
    1: ("blå", {"dark": "0038A8", "accent": "1F8A93", "light": "E3EAF7", "highlight": "FE4C61",
               "pale_text": "D6E0F5", "on_bg": "FFFFFF", "badge": "FE4C61"}),
    2: ("grøn", {"dark": "4F5400", "accent": "747B00", "light": "F1F4D4", "highlight": "0038A8",
                "pale_text": "EEF2C6", "on_bg": "333333", "badge": "0038A8"}),
    3: ("rød", {"dark": "A0142B", "accent": "D93349", "light": "FFE6DE", "highlight": "0038A8",
               "pale_text": "FFD9DF", "on_bg": "FFFFFF", "badge": "0038A8"}),
    4: ("creme", {"dark": "0038A8", "accent": "1F8A93", "light": "F8EDCA", "highlight": "FE4C61",
                 "pale_text": "D6E0F5", "on_bg": "333333", "badge": "FE4C61"}),
}


def _open_template() -> Presentation:
    """Åbn .potx som en præsentation, og fjern skabelonens eksempel-slides."""
    buffer = io.BytesIO()
    with (zipfile.ZipFile(TEMPLATE) as source,
          zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target):
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"presentationml.template.main+xml",
                                    b"presentationml.presentation.main+xml")
            target.writestr(item, data)
    buffer.seek(0)
    presentation = Presentation(buffer)
    slide_ids = presentation.slides._sldIdLst
    for slide_id in list(slide_ids):
        presentation.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)
    return presentation


prs = _open_template()
_family = FAMILIES["blå"]
_theme = {**COMMON, **PARTS[1][1]}


def use_part(number: int) -> None:
    """Vælg farvefamilie og farver for del 1-4. Kaldes først i main()."""
    global _family, _theme
    family_name, colors = PARTS[number]
    _family = FAMILIES[family_name]
    _theme = {**COMMON, **colors}


def save(out: Path, document_title: str) -> None:
    """Fjern ubrugte layouts (og deres fotos), og gem."""
    prs.core_properties.title = document_title
    used = {slide.slide_layout.name for slide in prs.slides}
    for layout in list(prs.slide_layouts):
        if layout.name not in used:
            prs.slide_layouts.remove(layout)
    prs.save(out)


# --- Enheder og farver -------------------------------------------------------------


def inch(value: float) -> int:
    return Inches(value * SCALE)


def pt(value: float) -> Pt:
    return Pt(value * SCALE)


def rgb(color: str) -> RGBColor:
    return RGBColor.from_string(_theme.get(color, color))


def _layout(name: str):
    return prs.slide_layouts.get_by_name(name)


def _drop_placeholders(slide, keep: tuple[int, ...] = ()) -> None:
    for placeholder in list(slide.placeholders):
        if placeholder.placeholder_format.idx not in keep:
            placeholder._element.getparent().remove(placeholder._element)


def _fill_placeholder(placeholder, content: str, x, y, w, h, size=None, anchor=MSO_ANCHOR.TOP):
    placeholder.left, placeholder.top, placeholder.width, placeholder.height = (
        inch(x), inch(y), inch(w), inch(h))
    frame = placeholder.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.text = content
    if size:
        for paragraph in frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = pt(size)


# --- Slides --------------------------------------------------------------------------


def new_slide():
    """En indholdsslide: familiens baggrund og logo. title() lægger panel og titel på."""
    slide = prs.slides.add_slide(_layout(_family.content))
    _drop_placeholders(slide, keep=(0,))
    return slide


def title(slide, heading, sub=None):
    """Titel på baggrunden, undertitel under den og et hvidt panel til indholdet."""
    top = 1.22 if sub else 1.05
    shape(slide, MSO_SHAPE.RECTANGLE, 0.3, top, 9.4, 5.08 - top, WHITE)
    _fill_placeholder(slide.shapes.title, heading, 0.5, 0.2, 9, 0.55, size=28)
    if sub:
        text(slide, 0.5, 0.78, 9, 0.35, sub, size=14, italic=True, color="on_bg")


def section(number, heading, sub):
    """Sektionsslide med skabelonens EK-mønster."""
    slide = prs.slides.add_slide(_layout(_family.section))
    badge(slide, 0.6, 1.2, str(number), "badge", size=0.7, font_size=24)
    _fill_placeholder(slide.placeholders[0], heading, 0.6, 2.05, 8.8, 0.9, size=36,
                      anchor=MSO_ANCHOR.BOTTOM)
    _fill_placeholder(slide.placeholders[1], sub, 0.6, 3.05, 8.8, 0.6, size=18)
    return slide


def title_slide(kicker, heading, sub, tags, note=None):
    """Forside med familiens foto til venstre og tekst til højre."""
    slide = prs.slides.add_slide(_layout(_family.title))
    text(slide, 4.6, 0.55, 5.0, 0.35, kicker, size=15, bold=True, color="on_bg")
    _fill_placeholder(slide.placeholders[0], heading, 4.6, 0.95, 5.0, 1.5, size=32,
                      anchor=MSO_ANCHOR.BOTTOM)
    _fill_placeholder(slide.placeholders[1], sub, 4.6, 2.6, 5.0, 1.0, size=16)
    for i, tag in enumerate(tags):
        x = 4.6 + (i % 2) * 2.5
        y = 3.75 + (i // 2) * 0.55
        shape(slide, MSO_SHAPE.RECTANGLE, x, y, 2.35, 0.45, WHITE)
        text(slide, x, y, 2.35, 0.45, tag, size=13, bold=True, color=DARK,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if note:
        notes(slide, note)
    return slide


def roadmap(heading, steps, lines):
    """Afslutning: trinene for undervisningen som kort, og tips under dem.

    steps: [(nummer, overskrift, undertekst, minutter)]. lines: linjer til text().
    """
    slide = new_slide()
    title(slide, heading)
    for i, (number, step_heading, sub, minutes) in enumerate(steps):
        x = 0.5 + i * 2.3
        card(slide, x, 1.3, 2.1, 1.8)
        badge(slide, x + 0.2, 1.5, number, CORAL if i == len(steps) - 1 else TEAL)
        text(slide, x + 0.2, 2.1, 1.8, 0.35, step_heading, size=15, bold=True, color=DARK)
        text(slide, x + 0.2, 2.42, 1.8, 0.3, sub, size=12, color=INK)
        text(slide, x + 0.2, 2.7, 1.8, 0.3, f"{minutes} min", size=12, color=MUTED)
    text(slide, 0.5, 3.45, 9, 1.5, lines, size=14, space_after=10)
    return slide


def notes(slide, content):
    slide.notes_slide.notes_text_frame.text = content


# --- Former og tekst ------------------------------------------------------------------


def text(slide, x, y, w, h, content, *, font=BODY, size=14, color=INK, bold=False,
         italic=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=0):
    """Tilføj en tekstboks.

    content er en streng (\\n = nyt afsnit), eller en liste af linjer, hvor hver linje
    er en liste af runs: (text, {"font", "size", "bold", "italic", "color", "link"}).
    """
    box = slide.shapes.add_textbox(inch(x), inch(y), inch(w), inch(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = anchor
    lines = [[(line, {})] for line in content.split("\n")] if isinstance(content, str) else content
    for i, runs in enumerate(lines):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_after = pt(space_after)
        for run_text, opts in runs:
            run = paragraph.add_run()
            run.text = run_text
            run.font.name = opts.get("font", font)
            factor = HEAD_SIZE_FACTOR if run.font.name == HEAD else 1
            run.font.size = pt(opts.get("size", size) * factor)
            run.font.bold = opts.get("bold", bold)
            run.font.italic = opts.get("italic", italic)
            run.font.color.rgb = rgb(opts.get("color", color))
            if "link" in opts:
                run.hyperlink.address = opts["link"]
    return box


def shape(slide, kind, x, y, w, h, fill, line=None, line_width=None, radius=0.1):
    item = slide.shapes.add_shape(kind, inch(x), inch(y), inch(w), inch(h))
    item.shadow.inherit = False
    item.fill.solid()
    item.fill.fore_color.rgb = rgb(fill)
    if line:
        item.line.color.rgb = rgb(line)
        if line_width:
            item.line.width = pt(line_width)
    else:
        item.line.fill.background()
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        item.adjustments[0] = min(0.5, radius / min(w, h))
    return item


def line(slide, x1, y1, x2, y2, color=DARK, width=1.75):
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, inch(x1), inch(y1), inch(x2), inch(y2)
    )
    connector.line.color.rgb = rgb(color)
    connector.line.width = pt(width)


def card(slide, x, y, w, h, fill=MINT):
    # Skabelonen bruger skarpe hjørner på sine bokse, så det gør vi også.
    return shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, fill)


def badge(slide, x, y, label, fill=TEAL, size=0.45, font_size=14):
    shape(slide, MSO_SHAPE.OVAL, x, y, size, size, fill)
    text(slide, x, y, size, size, label, size=font_size, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def code_box(slide, x, y, w, h, code, size=12):
    card(slide, x, y, w, h, CODE)
    text(slide, x + 0.15, y + 0.12, w - 0.3, h - 0.2, code, font=MONO, size=size)


def hint_cards(slide, x, y, w, items, card_h=0.95, gap=0.12, body_size=12):
    """Stak af kort, hver med en fed overskrift og en kort brødtekst."""
    for i, (heading, body) in enumerate(items):
        top = y + i * (card_h + gap)
        card(slide, x, top, w, card_h)
        text(slide, x + 0.15, top + 0.1, w - 0.3, 0.3, heading, font=HEAD, size=14,
             bold=True, color=DARK)
        text(slide, x + 0.15, top + 0.42, w - 0.3, card_h - 0.48, body, size=body_size)


def table(slide, x, y, widths, rows, *, size=12, row_h=0.36, mono_cols=(), first_fill=DARK):
    shape_ = slide.shapes.add_table(len(rows), len(widths), inch(x), inch(y),
                                    inch(sum(widths)), inch(row_h * len(rows)))
    grid = shape_.table
    for col, width in enumerate(widths):
        grid.columns[col].width = inch(width)
    for r, values in enumerate(rows):
        grid.rows[r].height = inch(row_h)
        for c, value in enumerate(values):
            cell = grid.cell(r, c)
            cell.margin_left = cell.margin_right = inch(0.08)
            cell.margin_top = cell.margin_bottom = inch(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            header = r == 0
            row_fill = first_fill if header else (MINT if r % 2 == 0 else WHITE)
            cell.fill.fore_color.rgb = rgb(row_fill)
            run = cell.text_frame.paragraphs[0].add_run()
            run.text = value
            run.font.name = MONO if (c in mono_cols and not header) else BODY
            run.font.size = pt(size)
            run.font.bold = header
            run.font.color.rgb = rgb(WHITE if header else INK)
    return grid


def warning(slide, y, label, body, h=0.55):
    card(slide, 0.5, y, 9, h, WARN_BG)
    text(slide, 0.7, y, 8.6, h, [[(label, {"bold": True, "color": BAD}), (body, {})]],
         size=14, anchor=MSO_ANCHOR.MIDDLE)


def arrow(slide, x, y, color=CORAL):
    shape(slide, MSO_SHAPE.RIGHT_ARROW, x, y, 0.3, 0.28, color)


# --- ER-figurer (del 2 og 3) ----------------------------------------------------------


def entity(slide, x, y, w, name, columns, row_h=0.33):
    """Tegn en ER-tabel: mørk overskrift og én række per (key, kolonne, type)."""
    shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, 0.45, DARK)
    text(slide, x, y, w, 0.45, name, font=HEAD, size=15, bold=True, color=WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    body_h = len(columns) * row_h + 0.1
    shape(slide, MSO_SHAPE.RECTANGLE, x, y + 0.45, w, body_h, WHITE, line=DARK, line_width=1.25)
    for i, (key, column, data_type) in enumerate(columns):
        top = y + 0.5 + i * row_h
        middle = {"anchor": MSO_ANCHOR.MIDDLE, "font": MONO}
        text(slide, x + 0.08, top, 0.4, row_h, key, size=10, bold=True, color=CORAL, **middle)
        text(slide, x + 0.45, top, w - 1.1, row_h, column, size=10, **middle)
        text(slide, x + w - 0.62, top, 0.6, row_h, data_type, size=8, color=MUTED, **middle)


def one_to_many(slide, parent_x, child_x, y):
    """Crow's foot-linje: || ved parent (præcis én), o< ved child (nul eller flere)."""
    line(slide, parent_x, y, child_x, y)
    for offset in (0.13, 0.23):
        line(slide, parent_x + offset, y - 0.12, parent_x + offset, y + 0.12)
    line(slide, child_x - 0.25, y, child_x, y - 0.13)
    line(slide, child_x - 0.25, y, child_x, y + 0.13)
    shape(slide, MSO_SHAPE.OVAL, child_x - 0.45, y - 0.07, 0.14, 0.14, WHITE, line=DARK,
          line_width=1.5)
